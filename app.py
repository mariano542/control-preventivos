from flask import Flask, render_template, request, session, redirect, url_for, send_file
import pandas as pd
from datetime import date
from io import BytesIO
import os
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

app = Flask(__name__)
app.secret_key = "apesa_tvm2026_secret"

PASSWORD = "TVM2026"
MARGEN_MATERIALES = 180
DIAS_PROXIMO = 45

def calcular_proximo(dias):
    if pd.isna(dias): return ""
    elif dias < 0: return "REVISAR"
    elif dias >= DIAS_PROXIMO: return "OK"
    else: return "PRÓXIMO"

def procesar_csv(archivo):
    try:
        df = pd.read_csv(archivo, sep=";", encoding="utf-8-sig")
        if len(df.columns) < 5:
            archivo.seek(0)
            df = pd.read_csv(archivo, sep=",", encoding="utf-8-sig")
    except:
        archivo.seek(0)
        df = pd.read_csv(archivo, sep=",", encoding="utf-8-sig")

    if "Activo" in df.columns:
        df = df.dropna(subset=["Activo"])
        df = df[df["Activo"].astype(str).str.strip() != ""]
        df = df[df["Activo"].astype(str).str.strip() != "Activo"]

    df["Frecuencia"] = pd.to_numeric(df.get("Frecuencia", pd.Series()), errors="coerce")
    df["Frecuencia acumulada"] = pd.to_numeric(df.get("Frecuencia acumulada", pd.Series()), errors="coerce")
    df["Fecha planificada"] = pd.to_datetime(df.get("Fecha planificada", pd.Series()), dayfirst=True, errors="coerce")

    fecha_hoy = pd.Timestamp(date.today())
    df["DIFERENCIA FRECUENCIA"] = df["Frecuencia"] - df["Frecuencia acumulada"]
    df["PEDIR MATERIAL"] = df["DIFERENCIA FRECUENCIA"].apply(
        lambda x: "PEDIR MATERIAL" if pd.notna(x) and x <= MARGEN_MATERIALES else "NO"
    )
    df["FECHA DE PARTIDA"] = date.today()
    df["fecha_solo"] = df["Fecha planificada"].dt.normalize()
    df["DIAS PARA TAREA"] = (df["fecha_solo"] - fecha_hoy).dt.days
    df["PRÓXIMO"] = df["DIAS PARA TAREA"].apply(calcular_proximo)
    df = df.drop(columns=["fecha_solo"], errors="ignore")
    return df.sort_values(["Tipo de Activo", "Activo"]).reset_index(drop=True)

def generar_excel(df):
    columnas = ["Tipo de mantenimiento","Tipo de Activo","Layout (Activo)","Activo",
                "Código de Frecuencia","Parte","Tarea","Fecha planificada","FECHA DE PARTIDA",
                "DIAS PARA TAREA","PRÓXIMO","Frecuencia","Frecuencia acumulada",
                "DIFERENCIA FRECUENCIA","PEDIR MATERIAL","Estado","OT"]
    cols = [c for c in columnas if c in df.columns]
    df_e = df[cols]
    wb = Workbook()
    wb.remove(wb.active)
    ln = Side(style="thin", color="CCCCCC")
    brd = Border(left=ln, right=ln, top=ln, bottom=ln)
    aw = {"Tipo de mantenimiento":18,"Tipo de Activo":22,"Layout (Activo)":35,"Activo":40,
          "Código de Frecuencia":14,"Parte":32,"Tarea":40,"Fecha planificada":18,
          "FECHA DE PARTIDA":16,"DIAS PARA TAREA":13,"PRÓXIMO":12,"Frecuencia":11,
          "Frecuencia acumulada":14,"DIFERENCIA FRECUENCIA":14,"PEDIR MATERIAL":15,"Estado":18,"OT":8}
    nw = {"DIFERENCIA FRECUENCIA","PEDIR MATERIAL","DIAS PARA TAREA","PRÓXIMO","FECHA DE PARTIDA"}

    def escribir(ws, datos):
        ws.freeze_panes = "A2"
        for ci, cn in enumerate(cols, 1):
            c = ws.cell(row=1, column=ci, value=cn)
            c.fill = PatternFill("solid", fgColor="C8102E" if cn in nw else "1A1A1A")
            c.font = Font(bold=True, color="FFFFFF", name="Arial", size=10)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = brd
            ws.column_dimensions[get_column_letter(ci)].width = aw.get(cn, 14)
        ws.row_dimensions[1].height = 30
        for ri, (_, fila) in enumerate(datos.iterrows(), 2):
            bf = PatternFill("solid", fgColor="F5F5F5" if ri%2==0 else "FFFFFF")
            for ci, cn in enumerate(cols, 1):
                val = fila[cn]
                c = ws.cell(row=ri, column=ci)
                if cn in ("Fecha planificada","FECHA DE PARTIDA"):
                    if pd.notna(val):
                        c.value = val.date() if hasattr(val,"date") else val
                        c.number_format = "DD/MM/YYYY"
                elif cn == "OT": c.value = int(val) if pd.notna(val) else None
                elif cn == "DIAS PARA TAREA": c.value = int(val) if pd.notna(val) else None
                else: c.value = val if pd.notna(val) else None
                c.font = Font(name="Arial", size=9, color="1A1A1A")
                c.border = brd
                if cn == "PEDIR MATERIAL":
                    c.fill = PatternFill("solid", fgColor="FFE699" if val=="PEDIR MATERIAL" else "E2EFDA")
                    c.alignment = Alignment(horizontal="center", vertical="center")
                    c.font = Font(name="Arial", size=9, bold=True)
                elif cn == "PRÓXIMO":
                    col = "FFCCCC" if val=="REVISAR" else ("E2EFDA" if val=="OK" else "FFE699")
                    c.fill = PatternFill("solid", fgColor=col)
                    c.alignment = Alignment(horizontal="center", vertical="center")
                    c.font = Font(name="Arial", size=9, bold=True)
                elif cn == "DIFERENCIA FRECUENCIA":
                    c.fill = PatternFill("solid", fgColor="FFE699" if pd.notna(val) and val<=MARGEN_MATERIALES else "E2EFDA")
                    c.alignment = Alignment(horizontal="center", vertical="center")
                elif cn == "DIAS PARA TAREA":
                    if pd.notna(val):
                        col = "FFCCCC" if val<0 else ("E2EFDA" if val>=DIAS_PROXIMO else "FFE699")
                        c.fill = PatternFill("solid", fgColor=col)
                    c.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    c.fill = bf
                    c.alignment = Alignment(horizontal="center" if ci in (1,2,5,12,13,16,17) else "left", vertical="center")
            ws.row_dimensions[ri].height = 16

    ws_all = wb.create_sheet("TODOS")
    escribir(ws_all, df_e)
    for tipo in sorted(df_e["Tipo de Activo"].dropna().unique()):
        ws = wb.create_sheet(tipo[:31])
        escribir(ws, df_e[df_e["Tipo de Activo"]==tipo])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf

@app.route("/", methods=["GET", "POST"])
def login():
    if session.get("autenticado"):
        return redirect(url_for("dashboard"))
    error = None
    if request.method == "POST":
        if request.form.get("password") == PASSWORD:
            session["autenticado"] = True
            return redirect(url_for("dashboard"))
        else:
            error = "Contraseña incorrecta"
    return render_template("login.html", error=error)

@app.route("/dashboard", methods=["GET", "POST"])
def dashboard():
    if not session.get("autenticado"):
        return redirect(url_for("login"))

    datos = None
    error = None
    stats = None
    fecha_hoy = date.today().strftime("%d/%m/%Y")

    if request.method == "POST":
        archivo = request.files.get("csv_file")
        if archivo and archivo.filename:
            try:
                df = procesar_csv(archivo)
                session["df_json"] = df.to_json(date_format="iso")

                stats = {
                    "total": len(df),
                    "pedir": len(df[df["PEDIR MATERIAL"]=="PEDIR MATERIAL"]),
                    "revisar": len(df[df["PRÓXIMO"]=="REVISAR"]),
                    "proximo": len(df[df["PRÓXIMO"]=="PRÓXIMO"]),
                    "ok": len(df[df["PRÓXIMO"]=="OK"]),
                }

                # Prepare table data
                cols_ver = ["Tipo de Activo","Activo","Parte","Tarea","Fecha planificada",
                           "DIAS PARA TAREA","PRÓXIMO","Frecuencia","Frecuencia acumulada",
                           "DIFERENCIA FRECUENCIA","PEDIR MATERIAL","Estado"]
                cols_disp = [c for c in cols_ver if c in df.columns]
                df_show = df[cols_disp].copy()
                if "Fecha planificada" in df_show.columns:
                    df_show["Fecha planificada"] = pd.to_datetime(df_show["Fecha planificada"]).dt.strftime("%d/%m/%Y")

                datos = df_show.fillna("").to_dict("records")
                datos_cols = cols_disp

            except Exception as e:
                error = f"Error al procesar el archivo: {str(e)}"
        else:
            error = "Por favor seleccioná un archivo CSV"

    return render_template("dashboard.html",
                          datos=datos,
                          stats=stats,
                          error=error,
                          fecha_hoy=fecha_hoy,
                          margen=MARGEN_MATERIALES,
                          dias_proximo=DIAS_PROXIMO)

@app.route("/descargar", methods=["POST"])
def descargar():
    if not session.get("autenticado"):
        return redirect(url_for("login"))
    try:
        df_json = session.get("df_json")
        if not df_json:
            return redirect(url_for("dashboard"))
        df = pd.read_json(df_json)
        excel_buf = generar_excel(df)
        nombre = f"APESA_Preventivos_{date.today().strftime('%Y%m%d')}.xlsx"
        return send_file(excel_buf, download_name=nombre,
                        as_attachment=True,
                        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    except Exception as e:
        return redirect(url_for("dashboard"))

@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))

if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))
    app.run(host="0.0.0.0", port=port)
